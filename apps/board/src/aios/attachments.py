"""Estrazione testo dagli allegati della chat (inline, non persistiti).

Immagini e PDF vanno NATIVI a Claude (blocchi image/document — li "vede" da solo).
Word/Excel/CSV/testo NON sono supportati nativi → qui li convertiamo in testo, che
viene iniettato nel prompt. Le dipendenze pesanti (python-docx, openpyxl) sono
importate lazy e degradano con un messaggio se non installate.
"""
from __future__ import annotations

import base64
import io


def _decode(data_b64: str) -> bytes:
    try:
        return base64.b64decode(data_b64 or "", validate=False)
    except Exception:
        return b""


def extract_text(name: str, media_type: str, data_b64: str, cap: int = 20000) -> str:
    """Testo estratto da un allegato Office/CSV/testo. Vuoto se non estraibile."""
    mt = (media_type or "").lower()
    nm = (name or "").lower()
    raw = _decode(data_b64)
    if not raw:
        return ""
    try:
        if mt.endswith("wordprocessingml.document") or nm.endswith(".docx"):
            return _docx(raw)[:cap]
        if mt.endswith("spreadsheetml.sheet") or nm.endswith(".xlsx"):
            return _xlsx(raw)[:cap]
        if mt == "text/csv" or nm.endswith(".csv"):
            return raw.decode("utf-8", "replace")[:cap]
        if mt.startswith("text/") or nm.endswith((".txt", ".md", ".json", ".log", ".yml", ".yaml")):
            return raw.decode("utf-8", "replace")[:cap]
    except Exception as exc:
        return f"(impossibile leggere «{name}»: {str(exc)[:120]})"
    return ""


def _docx(raw: bytes) -> str:
    try:
        import docx  # python-docx
    except Exception:
        return "(python-docx non installato: .docx non leggibile)"
    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in getattr(d, "tables", []):
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx(raw: bytes) -> str:
    try:
        import openpyxl
    except Exception:
        return "(openpyxl non installato: .xlsx non leggibile)"
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets[:5]:
        out.append(f"[Foglio: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 300:
                out.append("… (troncato)")
                break
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)
